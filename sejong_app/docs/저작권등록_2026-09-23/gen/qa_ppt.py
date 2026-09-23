# -*- coding: utf-8 -*-
"""생성된 .pptx 를 되읽어 (1) 글상자 넘침 검사 (2) 미리보기 PNG 렌더.

LibreOffice 가 없는 환경이라 실제 렌더 대신 글꼴 실측(Pillow)으로 줄바꿈을 계산한다.
한글은 Courier New 에 없어 PowerPoint 가 맑은 고딕으로 대체하므로, 글자마다 해당 글꼴로 잰다.
"""
import os, sys, math
from pptx import Presentation
from pptx.util import Emu
from PIL import Image, ImageDraw, ImageFont

OUT = os.path.dirname(os.path.abspath(__file__))
COUR = r"C:\Windows\Fonts\cour.ttf"
MALG = r"C:\Windows\Fonts\malgun.ttf"
CALI = r"C:\Windows\Fonts\calibri.ttf"
SCALE = 100          # px per inch (렌더 배율)
_cache = {}


def font(path, pt, px_per_pt=SCALE / 72.0):
    key = (path, round(pt, 2))
    if key not in _cache:
        _cache[key] = ImageFont.truetype(path, max(1, int(round(pt * px_per_pt))))
    return _cache[key]


def _fp(name):
    n = (name or '').lower()
    if 'courier' in n:
        return COUR
    return CALI


def measure(text, name, pt):
    """문자열 폭(px). 한글·기호는 맑은 고딕으로 대체해 잰다."""
    lat, kor = font(_fp(name), pt), font(MALG, pt)
    w = 0.0
    for ch in text:
        f = lat if ord(ch) < 0x2010 else kor
        try:
            w += f.getlength(ch)
        except Exception:
            w += lat.size * 0.6
    return w


def check(path):
    prs = Presentation(path)
    sw, sh = prs.slide_width / 914400, prs.slide_height / 914400
    problems = []
    images = []

    for si, slide in enumerate(prs.slides, 1):
        img = Image.new('RGB', (int(sw * SCALE), int(sh * SCALE)), 'white')
        d = ImageDraw.Draw(img)
        for shp in slide.shapes:
            x, y = shp.left / 914400, shp.top / 914400
            w, h = shp.width / 914400, shp.height / 914400
            if shp.shape_type is not None and not shp.has_text_frame:
                d.rectangle([x * SCALE, y * SCALE, (x + w) * SCALE, (y + h) * SCALE],
                            fill='#F4F5F8', outline='#D6DBE4')
            if not shp.has_text_frame:
                continue
            if shp.fill.type is not None and shp.shape_type is not None and shp.name.startswith('Rect'):
                d.rectangle([x * SCALE, y * SCALE, (x + w) * SCALE, (y + h) * SCALE],
                            fill='#F4F5F8', outline='#D6DBE4')
            tf = shp.text_frame
            avail = (w - (tf.margin_left + tf.margin_right) / 914400) * SCALE
            cy = y * SCALE + tf.margin_top / 914400 * SCALE
            used = 0.0
            for p in tf.paragraphs:
                runs = [(r.text, r.font.name, (r.font.size.pt if r.font.size else 12)) for r in p.runs]
                if not runs:
                    continue
                pt = max(s for _, _, s in runs)
                lead = (p.line_spacing.pt if hasattr(p.line_spacing, 'pt') and p.line_spacing
                        else (pt * 1.2 * (p.line_spacing if isinstance(p.line_spacing, float) else 1)))
                sb = p.space_before.pt if p.space_before else 0
                sa = p.space_after.pt if p.space_after else 0
                wpx = sum(measure(t, n, s) for t, n, s in runs)
                nlines = max(1, math.ceil(wpx / avail - 1e-6)) if avail > 0 else 1
                used += sb + nlines * lead + sa
                cy += sb * SCALE / 72.0
                cx = x * SCALE + tf.margin_left / 914400 * SCALE
                for t, n, s in runs:
                    for ch in t:                       # 한글은 맑은 고딕으로 그린다(대체 글꼴)
                        f = font(_fp(n) if ord(ch) < 0x2010 else MALG, s)
                        d.text((cx, cy), ch, font=f, fill='#1A1A1A')
                        cx += measure(ch, n, s)
                cy += nlines * lead * SCALE / 72.0 + sa * SCALE / 72.0
            avail_h = (h - (tf.margin_top + tf.margin_bottom) / 914400) * 72
            if used > avail_h + 0.5:
                problems.append('slide%d  %-22s  넘침 %.1fpt (내용 %.1f > 상자 %.1f)'
                                % (si, shp.name, used - avail_h, used, avail_h))
            if x < 0.4 or y < 0.3 or x + w > sw - 0.4 or y + h > sh - 0.3:
                problems.append('slide%d  %-22s  여백 부족 (x=%.2f y=%.2f x2=%.2f y2=%.2f)'
                                % (si, shp.name, x, y, x + w, y + h))
        p = os.path.join(OUT, os.path.basename(path).replace('.pptx', '') + '_s%d.png' % si)
        img.save(p)
        images.append(p)
    return problems, images


if __name__ == '__main__':
    for f in sys.argv[1:]:
        probs, imgs = check(f)
        print('===', os.path.basename(f))
        print('\n'.join('  ' + p for p in probs) if probs else '  넘침·여백 문제 없음')
        for i in imgs:
            print('  render:', i)
